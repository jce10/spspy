"""Human-readable spreadsheet exports for SPANC calibration results."""

import csv
from pathlib import Path
from typing import TYPE_CHECKING, Any

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

if TYPE_CHECKING:
    from .Spanc import Peak, Spanc


HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
SECTION_FILL = PatternFill("solid", fgColor="A9CCE3")


def _with_extension(fileName: str, extension: str) -> str:
    path = Path(fileName)
    if path.suffix.lower() != extension:
        path = path.with_suffix(extension)
    return str(path)


def _total_position_error(peak: "Peak") -> float:
    return float(np.hypot(peak.positionErrStat, peak.positionErrSys))


def _reaction_details(spanc: "Spanc", rxnName: str) -> dict[str, Any]:
    rxn = spanc.reactions[rxnName]
    return {
        "reaction_id": rxnName,
        # "reaction_label": str(rxn),
        "target": str(rxn.targetMaterial),
        "beam_energy_MeV": float(rxn.params.beamEnergy),
        "sps_angle_deg": float(np.degrees(rxn.params.spsAngle)),
        "magnetic_field_kG": float(rxn.params.magneticField),
        "ejectile_charge_state": int(rxn.params.ejectile_qs),
    }


def _write_section(sheet, title: str, headers: list[str], rows: list[list[Any]], startRow: int) -> int:
    titleCell = sheet.cell(row=startRow, column=1, value=title)
    titleCell.font = Font(bold=True)
    titleCell.fill = SECTION_FILL

    headerRow = startRow + 1
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=headerRow, column=column, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL

    for rowOffset, values in enumerate(rows, start=1):
        for column, value in enumerate(values, start=1):
            sheet.cell(row=headerRow + rowOffset, column=column, value=value)

    return headerRow + max(len(rows), 1) + 2


def _fit_parameter_unit(index: int, leadingUnit: str) -> str:
    if index == 0:
        return leadingUnit
    if index == 1:
        return f"{leadingUnit}/mm"
    return f"{leadingUnit}/mm^{index}"


def _size_columns(sheet) -> None:
    for columnCells in sheet.columns:
        maxLength = max(
            (len(str(cell.value)) for cell in columnCells if cell.value is not None),
            default=0,
        )
        sheet.column_dimensions[columnCells[0].column_letter].width = min(
            max(maxLength + 2, 12), 45
        )


def export_spanc_workbook(spanc: "Spanc", fileName: str) -> str:
    """Export the Data Tables and physical rho-fit results to an XLSX file."""
    fileName = _with_extension(fileName, ".xlsx")
    spanc.ensure_energy_fitters()

    fitAvailable = spanc.has_current_rho_fit()
    if fitAvailable:
        spanc.calculate_outputs()

    workbook = Workbook()
    dataSheet = workbook.active
    dataSheet.title = "Data Tables"
    fitSheet = workbook.create_sheet("Fit Results")

    row = 1
    targetRows = []
    for targetName, target in spanc.targets.items():
        for layerIndex, layer in enumerate(target.layer_details, start=1):
            compound = "; ".join(
                f"Z={z}, A={a}, stoichiometry={stoichiometry}"
                for z, a, stoichiometry in layer.compound_list
            )
            targetRows.append(
                [targetName, layerIndex, float(layer.thickness), compound]
            )
    row = _write_section(
        dataSheet,
        "Targets",
        ["Target", "Layer", "Thickness (ug/cm^2)", "Compound"],
        targetRows,
        row,
    )

    reactionRows = []
    for rxnName in spanc.reactions:
        details = _reaction_details(spanc, rxnName)
        reactionRows.append(
            [
                details["reaction_id"],
                details["reaction_label"],
                details["target"],
                details["beam_energy_MeV"],
                details["sps_angle_deg"],
                details["magnetic_field_kG"],
                details["ejectile_charge_state"],
            ]
        )
    row = _write_section(
        dataSheet,
        "Reactions",
        [
            "Reaction ID",
            "Reaction",
            "Target",
            "Beam Energy (MeV)",
            "SPS Angle (deg)",
            "Magnetic Field (kG)",
            "Ejectile Charge State",
        ],
        reactionRows,
        row,
    )

    calibrationRows = [
        [
            int(peak.peakID),
            peak.rxnName,
            float(peak.position),
            float(peak.positionErrStat),
            float(peak.positionErrSys),
            _total_position_error(peak),
            float(peak.rho),
            float(peak.rhoErr),
            float(peak.excitation),
            float(peak.excitationErr),
        ]
        for peak in spanc.calibrations.values()
    ]
    row = _write_section(
        dataSheet,
        "Calibration Peaks",
        [
            "Peak ID",
            "Reaction ID",
            "Position (mm)",
            "Position Stat. Uncertainty (mm)",
            "Position Sys. Uncertainty (mm)",
            "Total Position Uncertainty (mm)",
            "Rho (cm)",
            "Rho Uncertainty (cm)",
            "Assigned Ex (MeV)",
            "Ex Uncertainty (MeV)",
        ],
        calibrationRows,
        row,
    )

    outputRows = [
        [
            int(peak.peakID),
            peak.rxnName,
            float(peak.position),
            float(peak.positionErrStat),
            float(peak.positionErrSys),
            _total_position_error(peak),
            float(peak.rho),
            float(peak.rhoErr),
            float(peak.excitation),
            float(peak.excitationErr),
            float(peak.positionFWHM),
            float(peak.positionFWHMErr),
            float(peak.excitationFWHM),
            float(peak.excitationFWHMErr),
        ]
        for peak in spanc.outputs.values()
    ]
    _write_section(
        dataSheet,
        "Output Peaks",
        [
            "Peak ID",
            "Reaction ID",
            "Position (mm)",
            "Position Stat. Uncertainty (mm)",
            "Position Sys. Uncertainty (mm)",
            "Total Position Uncertainty (mm)",
            "Rho (cm)",
            "Rho Uncertainty (cm)",
            "Ex (MeV)",
            "Ex Uncertainty (MeV)",
            "Position FWHM (mm)",
            "Position FWHM Uncertainty (mm)",
            "Ex FWHM (MeV)",
            "Ex FWHM Uncertainty (MeV)",
        ],
        outputRows,
        row,
    )

    fitSheet["A1"] = "Physical Rho Calibration Fit"
    fitSheet["A1"].font = Font(bold=True)
    fitSheet["A1"].fill = SECTION_FILL
    fitSheet.append(["Fit Available", "Yes" if fitAvailable else "No"])

    if fitAvailable:
        fitter = spanc.fitter
        fitSheet.append(["Polynomial Order", int(fitter.polynomialOrder)])
        fitSheet.append(["Chi-Square", float(fitter.get_chisquare())])
        fitSheet.append(["NDF", int(fitter.get_ndf())])
        fitSheet.append(
            ["Reduced Chi-Square", float(fitter.get_reduced_chisquare())]
        )
        fitSheet.append(
            [
                "Equation",
                f"rho(cm) = {spanc.format_polynomial(fitter.get_parameters())}",
            ]
        )

        parameterStart = fitSheet.max_row + 3
        _write_section(
            fitSheet,
            "Fit Parameters",
            ["Parameter", "Value", "Uncertainty", "Units"],
            [
                [
                    f"a{index}",
                    float(value),
                    float(fitter.get_parameter_errors()[index]),
                    _fit_parameter_unit(index, "cm"),
                ]
                for index, value in enumerate(fitter.get_parameters())
            ],
            parameterStart,
        )

        residuals = spanc.get_residuals()
        energyResiduals, energyResidualErrors = spanc.get_energy_residuals()
        residualRows = []
        for peak, residual, energyResidual, energyResidualError in zip(
            spanc.calibrations.values(),
            residuals,
            energyResiduals,
            energyResidualErrors,
        ):
            fittedRho = float(fitter.evaluate(peak.position))
            fittedEx = float(
                spanc.reactions[peak.rxnName].calculate_excitation(fittedRho)
            )
            residualRows.append(
                [
                    int(peak.peakID),
                    peak.rxnName,
                    float(peak.position),
                    _total_position_error(peak),
                    float(peak.rho),
                    fittedRho,
                    float(residual.residual),
                    float(residual.residualError),
                    float(residual.studentizedResidual),
                    float(peak.excitation),
                    fittedEx,
                    float(energyResidual),
                    float(energyResidualError),
                ]
            )

        _write_section(
            fitSheet,
            "Calibration Residuals",
            [
                "Peak ID",
                "Reaction ID",
                "Position (mm)",
                "Total Position Uncertainty (mm)",
                "Assigned Rho (cm)",
                "Fitted Rho (cm)",
                "Rho Residual (cm)",
                "Rho Residual Uncertainty (cm)",
                "Studentized Residual",
                "Assigned Ex (MeV)",
                "Ex from Fitted Rho (MeV)",
                "Energy-Equivalent Residual (keV)",
                "Energy-Equivalent Residual Uncertainty (keV)",
            ],
            residualRows,
            fitSheet.max_row + 3,
        )
    else:
        fitSheet.append(
            [
                "Status",
                "No valid physical rho calibration fit was available at export.",
            ]
        )

    _size_columns(dataSheet)
    _size_columns(fitSheet)
    workbook.save(fileName)
    return fileName


def energy_calibration_rows(spanc: "Spanc", rxnName: str) -> list[dict[str, Any]]:
    """Build tidy, self-contained rows for a selected direct energy fit."""
    if not spanc.has_energy_fit(rxnName):
        raise ValueError(f"Reaction {rxnName} does not have an energy fit.")

    fitter = spanc.get_energy_fitter(rxnName)
    details = _reaction_details(spanc, rxnName)
    equationMeV, equationKeV = spanc.get_energy_equations(rxnName)
    parameters = fitter.get_parameters()
    parameterErrors = fitter.get_parameter_errors()
    warnings = " ".join(getattr(fitter, "lastWarnings", []))

    rows = []
    for peak in spanc.get_calibrations_for_reaction(rxnName):
        fittedEx = float(fitter.evaluate(peak.position))
        row = {
            **details,
            "polynomial_order": int(fitter.polynomialOrder),
            "equation_MeV": equationMeV,
            "equation_keV": equationKeV,
            "r_squared": float(fitter.get_r_squared()),
            "chi_square": float(fitter.get_chisquare()),
            "ndf": int(fitter.get_ndf()),
            "reduced_chi_square": float(fitter.get_reduced_chisquare()),
            "fit_warnings": warnings,
            "peak_id": int(peak.peakID),
            "position_mm": float(peak.position),
            "position_error_stat_mm": float(peak.positionErrStat),
            "position_error_sys_mm": float(peak.positionErrSys),
            "position_error_total_mm": _total_position_error(peak),
            "Ex_assigned_MeV": float(peak.excitation),
            "Ex_error_MeV": float(peak.excitationErr),
            "Ex_fit_MeV": fittedEx,
            "energy_residual_keV": float(
                (peak.excitation - fittedEx) * 1000.0
            ),
        }
        for index, value in enumerate(parameters):
            row[f"b{index}"] = float(value)
            row[f"ub{index}"] = float(parameterErrors[index])
        rows.append(row)
    return rows


def export_energy_calibration_csv(
    spanc: "Spanc", rxnName: str, fileName: str
) -> str:
    """Export one reaction's direct position-to-energy calibration to CSV."""
    fileName = _with_extension(fileName, ".csv")
    rows = energy_calibration_rows(spanc, rxnName)
    if not rows:
        raise ValueError(f"Reaction {rxnName} has no calibration points.")

    with open(fileName, "w", encoding="utf-8", newline="") as outputFile:
        writer = csv.DictWriter(outputFile, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    return fileName
